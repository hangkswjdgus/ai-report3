from flask import Flask, render_template, request, jsonify, send_file
from openai import OpenAI
from dotenv import load_dotenv
import os
import json
import re
from io import BytesIO
from typing import Any, Dict, List, Tuple
from datetime import datetime

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics

load_dotenv()

app = Flask(__name__)

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
MODEL_NAME = os.getenv("OPENAI_MODEL", "gpt-4.1-mini")

if not OPENAI_API_KEY:
    raise ValueError("OPENAI_API_KEY 환경변수가 설정되지 않았습니다.")

client = OpenAI(api_key=OPENAI_API_KEY)

ALLOWED_TEMPLATES = {"report", "paper", "resume", "presentation"}
ALLOWED_LENGTHS = {"짧게", "보통", "길게"}
ALLOWED_MODES = {
    "natural",
    "draft_assist",
    "revision_assist",
    "evidence_boost",
    "personalized",
}
ALLOWED_REF_STYLES = {"default", "apa", "mla", "chicago"}


def normalize_value(value: Any, default: str = "") -> str:
    if value is None:
        return default
    return str(value).strip()


def sanitize_input(data: Dict[str, Any]) -> Dict[str, str]:
    template = normalize_value(data.get("template"), "report")
    if template not in ALLOWED_TEMPLATES:
        template = "report"

    length = normalize_value(data.get("length"), "보통")
    if length not in ALLOWED_LENGTHS:
        length = "보통"

    mode = normalize_value(data.get("mode"), "natural")
    if mode not in ALLOWED_MODES:
        mode = "natural"

    ref_style = normalize_value(data.get("reference_style"), "apa").lower()
    if ref_style not in ALLOWED_REF_STYLES:
        ref_style = "apa"

    topic = normalize_value(data.get("topic"), "일반적인 주제")
    tone = normalize_value(data.get("tone"), "자연스럽고 이해하기 쉽게")
    purpose = normalize_value(data.get("purpose"), "일반적인 설명")
    details = normalize_value(data.get("details"), "기본적인 내용")
    extra = normalize_value(data.get("extra"), "없음")
    user_perspective = normalize_value(data.get("user_perspective"), "")
    class_context = normalize_value(data.get("class_context"), "")
    personal_experience = normalize_value(data.get("personal_experience"), "")
    professor_style = normalize_value(data.get("professor_style"), "")
    major = normalize_value(data.get("major"), "")
    target_company = normalize_value(data.get("target_company"), "")

    intro_weight = normalize_value(data.get("intro_weight"), "보통")
    body_weight = normalize_value(data.get("body_weight"), "보통")
    conclusion_weight = normalize_value(data.get("conclusion_weight"), "보통")
    balance_view = normalize_value(data.get("balance_view"), "일반")
    include_my_opinion = normalize_value(data.get("include_my_opinion"), "아니오")
    critical_view = normalize_value(data.get("critical_view"), "아니오")

    return {
        "template": template,
        "topic": topic,
        "length": length,
        "tone": tone,
        "purpose": purpose,
        "details": details,
        "extra": extra,
        "mode": mode,
        "user_perspective": user_perspective,
        "class_context": class_context,
        "personal_experience": personal_experience,
        "professor_style": professor_style,
        "major": major,
        "target_company": target_company,
        "reference_style": ref_style,
        "intro_weight": intro_weight,
        "body_weight": body_weight,
        "conclusion_weight": conclusion_weight,
        "balance_view": balance_view,
        "include_my_opinion": include_my_opinion,
        "critical_view": critical_view,
    }


def get_length_rule(length: str) -> str:
    if length == "짧게":
        return "본문은 최소 1600자 이상으로 작성할 것."
    elif length == "길게":
        return "본문은 최소 3600자 이상으로 작성할 것."
    return "본문은 최소 2400자 이상으로 작성할 것."


def get_mode_rules(mode: str) -> str:
    if mode == "natural":
        return """
작성 모드: 자연스러운 문체 모드
- 전체 문장을 너무 기계적으로 끊지 말고 자연스럽게 이어질 것.
- 보고서 문체를 유지하되 지나치게 딱딱하거나 교과서식 문장만 반복하지 말 것.
- 근거를 제시한 뒤 그 의미를 사람 독자가 이해하기 쉽게 풀어 설명할 것.
"""
    if mode == "draft_assist":
        return """
작성 모드: 초안 보조 모드
- 완성본처럼 쓰되, 사용자가 이후에 수정하기 쉽도록 구조를 명확하게 제시할 것.
- 각 소제목 아래 핵심 논지를 분명히 두고, 과도한 미사여구보다 초안 작성 효율을 높일 것.
- 문단 흐름은 자연스럽게 유지하되 수정 가능한 여지를 남기는 정돈된 톤으로 작성할 것.
"""
    if mode == "revision_assist":
        return """
작성 모드: 수정 보조 모드
- 이미 사용자가 작성한 문서를 다듬는 것처럼 자연스럽고 정리된 문체를 사용할 것.
- 중복 표현을 줄이고, 문장 연결을 매끄럽게 만들 것.
- 정보 나열보다 문단 간 논리 연결을 강화할 것.
"""
    if mode == "evidence_boost":
        return """
작성 모드: 근거 강화 모드
- 핵심 주장마다 가능한 한 구체적인 통계, 보고서, 논문, 공공기관 자료를 붙일 것.
- 단, 자료 나열형이 되지 않도록 반드시 해석과 시사점을 덧붙일 것.
- 인용은 본문 흐름을 해치지 않게 자연스럽게 녹여낼 것.
"""
    if mode == "personalized":
        return """
작성 모드: 개인화 모드
- 사용자가 제공한 전공, 수업 맥락, 관점, 경험, 관심사, 교수 요구사항을 적극 반영할 것.
- 단, 사용자가 제공하지 않은 개인 경험이나 감정은 지어내지 말 것.
- 문체는 자연스럽고 개별 작성자 관점이 느껴지게 하되, AI 사용 사실을 숨기기 위한 목적의 표현 최적화는 하지 말 것.
- "나는", "본인은", "이번 수업에서", "내가 주목한 점은" 등의 표현은 사용자가 제공한 맥락이 있을 때만 제한적으로 활용할 것.
"""
    return ""


def get_reference_style_instruction(style: str) -> str:
    if style == "apa":
        return "참고문헌은 APA 스타일에 최대한 가깝게 정리할 것."
    if style == "mla":
        return "참고문헌은 MLA 스타일에 최대한 가깝게 정리할 것."
    if style == "chicago":
        return "참고문헌은 Chicago 스타일에 최대한 가깝게 정리할 것."
    return "참고문헌은 읽기 좋은 기본 학술 형식으로 정리할 것."


def get_assignment_rules(payload: Dict[str, str]) -> str:
    intro_weight = payload["intro_weight"]
    body_weight = payload["body_weight"]
    conclusion_weight = payload["conclusion_weight"]
    balance_view = payload["balance_view"]
    include_my_opinion = payload["include_my_opinion"]
    critical_view = payload["critical_view"]

    return f"""
과제 특화 옵션:
- 서론 분량 비중: {intro_weight}
- 본론 분량 비중: {body_weight}
- 결론 분량 비중: {conclusion_weight}
- 찬반 균형 설정: {balance_view}
- 내 의견 포함 여부: {include_my_opinion}
- 비판적 관점 포함 여부: {critical_view}

적용 규칙:
1. 분량 비중 설정을 실제 문단 길이에 반영할 것.
2. 찬반 균형이 "균형 있게"이면 장점과 한계를 모두 다룰 것.
3. 찬반 균형이 "찬성 중심"이면 긍정적 효과를 중심으로 쓰되 한계도 짧게 언급할 것.
4. 찬반 균형이 "비판 중심"이면 문제점과 한계를 중심으로 쓰되 필요 시 장점도 짧게 언급할 것.
5. 비판적 관점 포함이 "예"이면 근거 기반의 한계, 반론, 주의점도 함께 쓸 것.
6. 내 의견 포함이 "예"이면 결론 또는 본론 후반부에 사용자의 관점을 드러내는 문단을 포함할 수 있다. 단, 사용자가 제공하지 않은 경험은 지어내지 말 것.
"""


def get_template_structure(
    template: str,
    topic: str,
    purpose: str,
    details: str,
    extra: str,
    major: str,
    target_company: str,
    class_context: str,
    user_perspective: str,
    personal_experience: str,
    professor_style: str,
) -> str:
    if template == "report":
        return f"""
문서 유형: 대학 보고서

반드시 아래 구조를 본문(body)에 포함할 것:
# 서론
# 본론
# 결론

보고서 작성 원칙:
- 서론에서는 주제의 배경, 중요성, 문제의식을 충분히 설명할 것.
- 본론에서는 주장 → 근거 → 해석 → 시사점 흐름으로 서술할 것.
- 결론에서는 단순 요약을 넘어서 최종적인 판단과 방향성을 제시할 것.

사용자 입력:
주제: {topic}
작성 목적: {purpose}
포함하고 싶은 내용: {details}
추가 요구사항: {extra}
수업 맥락: {class_context}
사용자 관점: {user_perspective}
개인 경험 또는 생각: {personal_experience}
교수/과제 스타일: {professor_style}
전공: {major}
"""
    if template == "paper":
        return f"""
문서 유형: 논문 형식 초안

반드시 아래 구조를 본문(body)에 포함할 것:
# 초록
# 서론
# 관련 연구 또는 이론적 배경
# 본문 분석
# 결론

논문 형식 작성 원칙:
- 학술적이되 지나치게 부자연스럽지 않게 작성할 것.
- 관련 연구에서는 기존 연구의 공통점, 차이점, 한계도 간단히 언급할 것.
- 본문 분석에서는 자료 해석과 논리적 주장 전개를 분명히 할 것.
- 본문 내 인용은 학술 글처럼 자연스럽게 배치할 것.

사용자 입력:
주제: {topic}
연구 목적: {purpose}
연구 내용: {details}
추가 요구사항: {extra}
수업 맥락: {class_context}
사용자 관점: {user_perspective}
교수/과제 스타일: {professor_style}
전공: {major}
"""
    if template == "resume":
        return f"""
문서 유형: 자기소개서 / 이력서 보조문

반드시 아래 구조를 본문(body)에 포함할 것:
# 지원 동기
# 경험 및 역량
# 강점
# 마무리

자기소개서 작성 원칙:
- 단순 스펙 나열이 아니라 경험이 역량으로 이어지는 흐름을 만들 것.
- 사용자가 제공한 경험과 관점만 활용할 것.
- 제공되지 않은 수상, 인턴, 프로젝트, 자격증, 성과는 지어내지 말 것.
- 지원 기업 또는 지원 맥락이 있으면 그에 맞춰 표현을 조정할 것.

사용자 입력:
주제: {topic}
지원 목적: {purpose}
경험 및 내용: {details}
추가 요구사항: {extra}
개인 경험 또는 생각: {personal_experience}
사용자 관점: {user_perspective}
지원 대상 기업/기관: {target_company}
전공: {major}
"""
    return f"""
문서 유형: 발표문

반드시 아래 구조를 본문(body)에 포함할 것:
# 도입
# 핵심 내용
# 근거 및 사례
# 결론

발표문 작성 원칙:
- 발표용 문서지만 메모 수준으로 짧지 않게, 실제 발표 원고처럼 자연스럽게 작성할 것.
- 핵심 메시지가 분명히 드러나도록 할 것.
- 근거를 제시한 뒤 왜 중요한지까지 설명할 것.

사용자 입력:
주제: {topic}
발표 목적: {purpose}
핵심 내용: {details}
추가 요구사항: {extra}
수업 맥락: {class_context}
사용자 관점: {user_perspective}
개인 경험 또는 생각: {personal_experience}
전공: {major}
"""


def build_prompt(payload: Dict[str, str]) -> str:
    template = payload["template"]
    topic = payload["topic"]
    length = payload["length"]
    tone = payload["tone"]
    purpose = payload["purpose"]
    details = payload["details"]
    extra = payload["extra"]
    mode = payload["mode"]
    reference_style = payload["reference_style"]

    user_perspective = payload["user_perspective"]
    class_context = payload["class_context"]
    personal_experience = payload["personal_experience"]
    professor_style = payload["professor_style"]
    major = payload["major"]
    target_company = payload["target_company"]

    length_rule = get_length_rule(length)
    mode_rules = get_mode_rules(mode)
    ref_style_instruction = get_reference_style_instruction(reference_style)
    assignment_rules = get_assignment_rules(payload)

    common_rules = f"""
너는 상업용 문서 생성 서비스의 고급 작성 엔진이다.
사용자의 입력을 바탕으로 대학 과제, 보고서, 발표문, 논문 초안, 자기소개서 초안을 구조적으로 작성한다.

절대 규칙:
1. 출처가 불분명한 통계, 수치, 연구명, 기관명, 논문명, 저자, 퍼센트, 연도는 절대 지어내지 말 것.
2. 사용자가 제공하지 않은 개인 경험, 성과, 수상, 인턴, 감정, 발언은 지어내지 말 것.
3. 핵심 주장에는 가능한 한 실제 자료, 통계, 논문, 보고서, 공공기관 자료를 근거로 붙일 것.
4. 근거를 제시한 뒤에는 반드시 그 의미를 해석하고, 주제와 연결되는 주장까지 전개할 것.
5. 단순히 "연구에 따르면"으로 끝내지 말고, 가능하면 기관명, 저자, 연도, 수치 또는 핵심 결과를 함께 제시할 것.
6. 숫자나 통계가 있으면 그 수치가 왜 중요한지 설명할 것.
7. 수치가 없더라도 신뢰 가능한 기관명, 보고서명, 논문명을 기반으로 서술할 수 있다. 단, 출처 없는 내용은 금지한다.
8. 문단은 짧게 끝내지 말고, 주장 → 근거 → 해석 → 시사점 흐름으로 충분히 전개할 것.
9. 문단 간 연결 문장을 넣어 전체 글의 흐름이 자연스럽게 이어지도록 할 것.
10. 문체는 {tone}를 기본으로 하되, 사용자 목적과 문서 유형에 맞게 조정할 것.
11. {length_rule}
12. 마지막에는 references 배열에 실제 사용한 참고자료만 정리할 것.
13. {ref_style_instruction}
14. 본문에는 가능한 경우 인용 표시를 넣을 것.
15. 인용 표시는 자연스럽게 다음 형태를 우선 사용할 것:
   - (저자, 연도)
   - (기관명, 연도)
   - 문장형 예: 한국교육학술정보원(2023)에 따르면 ...
16. references는 본문에 실제로 반영된 자료만 넣을 것.
17. citations에는 본문에 반영한 주요 인용 근거를 짧게 요약해서 넣을 것.
18. JSON 객체만 출력할 것. 코드블록, 설명문, 서문은 금지한다.
19. body에는 문서 본문만 넣을 것.
20. output_language는 반드시 "ko"로 넣을 것.

반드시 아래 JSON 형식으로만 응답해라:
{{
  "title": "문서 제목",
  "body": "문서 본문 전체",
  "references": [
    "참고문헌 1",
    "참고문헌 2"
  ],
  "citations": [
    "본문에 사용한 핵심 인용 1",
    "본문에 사용한 핵심 인용 2"
  ],
  "output_language": "ko"
}}
"""

    structure = get_template_structure(
        template=template,
        topic=topic,
        purpose=purpose,
        details=details,
        extra=extra,
        major=major,
        target_company=target_company,
        class_context=class_context,
        user_perspective=user_perspective,
        personal_experience=personal_experience,
        professor_style=professor_style,
    )

    web_rules = """
특히 중요:
- 웹 검색을 통해 실제로 확인 가능한 자료를 우선 사용할 것.
- 한국 자료가 있으면 우선 사용하고, 부족하면 해외 공공기관, 대학, 학술 자료를 추가할 것.
- 정부기관, 공공기관, 대학, 학술지, 국제기구 자료를 우선 사용할 것.
- 수치가 포함된 자료가 있으면 우선 활용하되, 수치만 던지고 끝내지 말고 의미를 설명할 것.
- references에는 실제 사용한 자료만 넣을 것.
- 글 전체가 자료 나열문처럼 보이지 않도록 해석과 주장을 함께 쓸 것.
"""

    return f"{common_rules}\n{mode_rules}\n{assignment_rules}\n{structure}\n{web_rules}"


def build_refine_prompt(action: str, title: str, body: str, template: str, tone: str) -> str:
    action_map = {
        "polish_style": """
목표:
- 문체를 더 매끄럽고 자연스럽게 다듬을 것.
- 의미는 유지하되 문장 연결과 표현을 개선할 것.
- 지나치게 딱딱하거나 반복적인 표현을 줄일 것.
""",
        "strengthen_evidence": """
목표:
- 본문의 핵심 주장에 더 구체적인 근거와 자료를 보강할 것.
- 가능하면 신뢰 가능한 통계, 연구, 보고서, 기관 자료를 추가할 것.
- 단, 지어내지 말고 실제 확인 가능한 내용만 사용할 것.
""",
        "expand_conclusion": """
목표:
- 결론 부분을 더 풍부하게 확장할 것.
- 단순 요약을 넘어서 시사점, 한계, 향후 방향까지 포함할 것.
""",
        "presentation_summary": """
목표:
- 원문을 바탕으로 발표용 요약본 성격을 일부 반영해 문장을 더 명확하게 정리할 것.
- 핵심 메시지가 잘 드러나도록 문단을 정돈할 것.
- 원문 길이를 너무 심하게 줄이지는 말 것.
""",
        "add_critical_view": """
목표:
- 기존 글에 비판적 관점과 한계, 반론, 주의점을 추가할 것.
- 균형 잡힌 시각이 드러나도록 쓸 것.
"""
    }

    selected_action = action_map.get(action, action_map["polish_style"])

    return f"""
너는 사용자가 이미 생성한 문서를 후처리하는 편집 엔진이다.

문서 유형: {template}
기본 문체: {tone}

{selected_action}

절대 규칙:
1. 원문 구조를 최대한 유지할 것.
2. 제목은 유지하되 필요하면 더 자연스럽게 다듬을 수 있다.
3. 근거를 보강할 때는 실제 확인 가능한 내용만 사용할 것.
4. 지어낸 통계, 논문, 기관명을 넣지 말 것.
5. 결과는 JSON 형식으로만 출력할 것.

반드시 아래 JSON 형식으로만 응답:
{{
  "title": "수정된 제목",
  "body": "수정된 본문",
  "references": [
    "참고문헌 1"
  ],
  "citations": [
    "핵심 인용 1"
  ],
  "output_language": "ko"
}}

원문 제목:
{title}

원문 본문:
{body}
"""


def build_references_prompt(title: str, body: str, reference_style: str) -> str:
    ref_style_instruction = get_reference_style_instruction(reference_style)

    return f"""
너는 문서 본문을 읽고 참고문헌과 본문 인용 요약만 재정리하는 엔진이다.

절대 규칙:
1. 본문에 실제로 반영된 자료만 추정하여 정리할 것.
2. 확인할 수 없는 자료는 지어내지 말 것.
3. {ref_style_instruction}
4. 결과는 JSON 형식으로만 출력할 것.
5. body는 수정하지 말고 그대로 반환할 것.

반드시 아래 JSON 형식으로만 응답:
{{
  "title": "{title}",
  "body": "원문 본문 그대로",
  "references": [
    "참고문헌 1",
    "참고문헌 2"
  ],
  "citations": [
    "본문에서 사용된 핵심 인용 1",
    "본문에서 사용된 핵심 인용 2"
  ],
  "output_language": "ko"
}}

제목:
{title}

본문:
{body}
"""


def extract_json_text(raw_text: str) -> str:
    text = raw_text.strip()

    if text.startswith("```json"):
        text = text[7:]
    elif text.startswith("```"):
        text = text[3:]

    if text.endswith("```"):
        text = text[:-3]

    text = text.strip()

    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1 and end > start:
        text = text[start:end + 1]

    return text.strip()


def parse_model_response(raw_text: str) -> Tuple[str, str, List[str], List[str], str]:
    cleaned_text = extract_json_text(raw_text)

    try:
        parsed = json.loads(cleaned_text)

        title = normalize_value(parsed.get("title"), "생성된 문서")
        body = normalize_value(parsed.get("body"), "")
        references = parsed.get("references", [])
        citations = parsed.get("citations", [])
        output_language = normalize_value(parsed.get("output_language"), "ko")

        if not isinstance(references, list):
            references = [str(references)]

        if not isinstance(citations, list):
            citations = [str(citations)]

        references = [normalize_value(ref) for ref in references if normalize_value(ref)]
        citations = [normalize_value(c) for c in citations if normalize_value(c)]

        if not title:
            title = "생성된 문서"

        if not body:
            body = "본문을 받아오지 못했습니다."

        return title, body, references, citations, output_language

    except Exception:
        title = "생성된 문서"
        body = raw_text.strip()
        references = ["참고문헌 파싱에 실패했습니다. 결과를 직접 확인해 주세요."]
        citations = []
        output_language = "ko"

        try:
            title_match = re.search(r'"title"\s*:\s*"(.*?)"\s*,', raw_text, re.DOTALL)
            body_match = re.search(r'"body"\s*:\s*"(.*?)"\s*(,\s*"references"|,\s*"citations"|})', raw_text, re.DOTALL)
            refs_match = re.search(r'"references"\s*:\s*(\[[\s\S]*?\])', raw_text, re.DOTALL)
            citations_match = re.search(r'"citations"\s*:\s*(\[[\s\S]*?\])', raw_text, re.DOTALL)

            if title_match:
                title = title_match.group(1).replace("\\n", "\n").replace('\\"', '"').strip()

            if body_match:
                body = body_match.group(1).replace("\\n", "\n").replace('\\"', '"').strip()

            if refs_match:
                parsed_refs = json.loads(refs_match.group(1))
                if isinstance(parsed_refs, list):
                    references = [normalize_value(ref) for ref in parsed_refs if normalize_value(ref)]

            if citations_match:
                parsed_citations = json.loads(citations_match.group(1))
                if isinstance(parsed_citations, list):
                    citations = [normalize_value(c) for c in parsed_citations if normalize_value(c)]

        except Exception:
            pass

        return title, body, references, citations, output_language


def safe_filename(title: str, ext: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', "", title).strip()
    cleaned = cleaned.replace(" ", "_")
    if not cleaned:
        cleaned = f"document_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    return f"{cleaned}.{ext}"


def set_run_korean_font(run, font_name="Malgun Gothic", font_size=11, bold=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold

    r = run._element
    r.rPr.rFonts.set(qn("w:ascii"), font_name)
    r.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    r.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    r.rPr.rFonts.set(qn("w:cs"), font_name)


def create_docx_file(title: str, body: str, references: List[str], citations: List[str]) -> BytesIO:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Malgun Gothic")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Malgun Gothic")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    style._element.rPr.rFonts.set(qn("w:cs"), "Malgun Gothic")

    title_para = doc.add_paragraph()
    title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    title_run = title_para.add_run(title)
    set_run_korean_font(title_run, font_name="Malgun Gothic", font_size=16, bold=True)

    doc.add_paragraph("")

    for line in body.split("\n"):
        stripped = line.strip()

        if not stripped:
            doc.add_paragraph("")
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(stripped[2:].strip())
            set_run_korean_font(r, font_name="Malgun Gothic", font_size=14, bold=True)
        else:
            p = doc.add_paragraph()
            r = p.add_run(stripped)
            set_run_korean_font(r, font_name="Malgun Gothic", font_size=11, bold=False)

    if citations:
        doc.add_paragraph("")
        p = doc.add_paragraph()
        r = p.add_run("인용 요약")
        set_run_korean_font(r, font_name="Malgun Gothic", font_size=13, bold=True)

        for item in citations:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(item)
            set_run_korean_font(r, font_name="Malgun Gothic", font_size=11, bold=False)

    if references:
        doc.add_paragraph("")
        p = doc.add_paragraph()
        r = p.add_run("참고문헌")
        set_run_korean_font(r, font_name="Malgun Gothic", font_size=13, bold=True)

        for ref in references:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(ref)
            set_run_korean_font(r, font_name="Malgun Gothic", font_size=11, bold=False)

    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


def register_korean_font() -> str:
    font_path = os.path.join(app.root_path, "fonts", "NanumGothic.ttf")
    font_name = "Helvetica"

    if os.path.exists(font_path):
        pdfmetrics.registerFont(TTFont("NanumGothic", font_path))
        font_name = "NanumGothic"

    return font_name


def create_pdf_file(title: str, body: str, references: List[str], citations: List[str]) -> BytesIO:
    file_stream = BytesIO()
    font_name = register_korean_font()

    doc = SimpleDocTemplate(
        file_stream,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "KoreanTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=16,
        leading=22,
        spaceAfter=18
    )

    heading_style = ParagraphStyle(
        "KoreanHeading",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=13,
        leading=18,
        spaceBefore=12,
        spaceAfter=8
    )

    body_style = ParagraphStyle(
        "KoreanBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10.5,
        leading=16,
        spaceAfter=8
    )

    story = []
    story.append(Paragraph(title.replace("\n", "<br/>"), title_style))
    story.append(Spacer(1, 10))

    for line in body.split("\n"):
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 6))
            continue

        if stripped.startswith("# "):
            story.append(Paragraph(stripped[2:].strip(), heading_style))
        else:
            safe_line = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe_line, body_style))

    if citations:
        story.append(Spacer(1, 12))
        story.append(Paragraph("인용 요약", heading_style))
        for item in citations:
            safe_item = f"• {item}".replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe_item, body_style))

    if references:
        story.append(Spacer(1, 12))
        story.append(Paragraph("참고문헌", heading_style))
        for ref in references:
            safe_ref = f"• {ref}".replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe_ref, body_style))

    doc.build(story)
    file_stream.seek(0)
    return file_stream


@app.route("/")
def home():
    return render_template("index.html")


@app.route("/health", methods=["GET"])
def health():
    return jsonify({
        "status": "ok",
        "service": "ai-document-generator",
        "model": MODEL_NAME
    })


@app.route("/generate", methods=["POST"])
def generate():
    try:
        data = request.get_json(silent=True) or {}
        payload = sanitize_input(data)

        prompt = build_prompt(payload)

        response = client.responses.create(
            model=MODEL_NAME,
            tools=[{"type": "web_search_preview"}],
            input=prompt
        )

        raw_text = response.output_text.strip()
        title, body, references, citations, output_language = parse_model_response(raw_text)

        return jsonify({
            "success": True,
            "title": title,
            "body": body,
            "references": references,
            "citations": citations,
            "output_language": output_language,
            "meta": {
                "template": payload["template"],
                "mode": payload["mode"],
                "length": payload["length"],
                "tone": payload["tone"],
                "reference_style": payload["reference_style"]
            }
        })

    except Exception as e:
        return jsonify({
            "success": False,
            "title": "오류",
            "body": f"오류 발생: {str(e)}",
            "references": [],
            "citations": [],
            "output_language": "ko"
        }), 500


@app.route("/refine", methods=["POST"])
def refine():
    try:
        data = request.get_json(silent=True) or {}

        action = normalize_value(data.get("action"), "polish_style")
        title = normalize_value(data.get("title"), "생성된 문서")
        body = normalize_value(data.get("body"), "")
        template = normalize_value(data.get("template"), "report")
        tone = normalize_value(data.get("tone"), "자연스럽고 이해하기 쉽게")

        prompt = build_refine_prompt(action, title, body, template, tone)

        response = client.responses.create(
            model=MODEL_NAME,
            tools=[{"type": "web_search_preview"}],
            input=prompt
        )

        raw_text = response.output_text.strip()
        new_title, new_body, references, citations, output_language = parse_model_response(raw_text)

        return jsonify({
            "success": True,
            "title": new_title,
            "body": new_body,
            "references": references,
            "citations": citations,
            "output_language": output_language
        })

    except Exception as e:
        return jsonify({
            "success": False,
            "message": f"후처리 중 오류 발생: {str(e)}"
        }), 500


@app.route("/rebuild_references", methods=["POST"])
def rebuild_references():
    try:
        data = request.get_json(silent=True) or {}

        title = normalize_value(data.get("title"), "생성된 문서")
        body = normalize_value(data.get("body"), "")
        reference_style = normalize_value(data.get("reference_style"), "apa").lower()
        if reference_style not in ALLOWED_REF_STYLES:
            reference_style = "apa"

        prompt = build_references_prompt(title, body, reference_style)

        response = client.responses.create(
            model=MODEL_NAME,
            tools=[{"type": "web_search_preview"}],
            input=prompt
        )

        raw_text = response.output_text.strip()
        new_title, new_body, references, citations, output_language = parse_model_response(raw_text)

        return jsonify({
            "success": True,
            "title": new_title,
            "body": new_body,
            "references": references,
            "citations": citations,
            "output_language": output_language
        })

    except Exception as e:
        return jsonify({
            "success": False,
            "message": f"참고문헌 재생성 중 오류 발생: {str(e)}"
        }), 500


@app.route("/download/docx", methods=["POST"])
def download_docx():
    try:
        data = request.get_json(silent=True) or {}

        title = normalize_value(data.get("title"), "생성된 문서")
        body = normalize_value(data.get("body"), "")
        references = data.get("references", [])
        citations = data.get("citations", [])

        if not isinstance(references, list):
            references = []
        if not isinstance(citations, list):
            citations = []

        file_stream = create_docx_file(title, body, references, citations)
        filename = safe_filename(title, "docx")

        return send_file(
            file_stream,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    except Exception as e:
        return jsonify({
            "success": False,
            "message": f"DOCX 생성 중 오류 발생: {str(e)}"
        }), 500


@app.route("/download/pdf", methods=["POST"])
def download_pdf():
    try:
        data = request.get_json(silent=True) or {}

        title = normalize_value(data.get("title"), "생성된 문서")
        body = normalize_value(data.get("body"), "")
        references = data.get("references", [])
        citations = data.get("citations", [])

        if not isinstance(references, list):
            references = []
        if not isinstance(citations, list):
            citations = []

        file_stream = create_pdf_file(title, body, references, citations)
        filename = safe_filename(title, "pdf")

        return send_file(
            file_stream,
            as_attachment=True,
            download_name=filename,
            mimetype="application/pdf"
        )

    except Exception as e:
        return jsonify({
            "success": False,
            "message": f"PDF 생성 중 오류 발생: {str(e)}"
        }), 500


if __name__ == "__main__":
    port = int(os.getenv("PORT", 5000))
    app.run(host="0.0.0.0", port=port)